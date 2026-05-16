#!/usr/bin/env python3
"""Collect RULER lm-eval results into tidy CSV summaries.

Expected layout:
  <input_root>/<model_dir>/<task>-<setting>.json

Outputs:
  - ruler_summary_one_table.csv
      one row per (model, setting), with grouped two-row headers:
      niah / qa / vt / cwe / fwe / overall
      and sub-columns 4k / 8k / ... / avg
  - ruler_summary_one_table.xlsx
      same grouped table, but rendered properly with merged multi-row headers
  - ruler_results_long.csv
      one row per (model, setting, category, task, seq_len)
  - ruler_task_by_length.csv
      one row per (model, setting, category, task), lengths as columns
  - ruler_category_by_length.csv
      one row per (model, setting, category), averaged over tasks in category
  - ruler_overall_by_length.csv
      one row per (model, setting), averaged over all subtasks
"""
from __future__ import annotations

import json
import os
import re
from pathlib import Path
from typing import Any

import pandas as pd


_TIMESTAMP_RE = re.compile(r"_\d{4}-\d{2}-\d{2}T[\d\-\.]+$")
_LEN_METRIC_RE = re.compile(r"^(\d+),none$")

SEQ_LENGTH_ORDER = [4096, 8192, 16384, 32768, 65536, 131072]
CATEGORY_ORDER = {
    "niah": 0,
    "qa": 1,
    "vt": 2,
    "cwe": 3,
    "fwe": 4,
    "other": 5,
    "overall": 6,
}


def strip_timestamp_suffix(stem: str) -> str:
    return _TIMESTAMP_RE.sub("", stem)


def categorize_task(task_name: str) -> str:
    if task_name.startswith("niah_"):
        return "niah"
    if task_name.startswith("ruler_qa_"):
        return "qa"
    if task_name == "ruler_vt":
        return "vt"
    if task_name == "ruler_cwe":
        return "cwe"
    if task_name == "ruler_fwe":
        return "fwe"
    if task_name == "ruler":
        return "overall"
    return "other"


def parse_task_setting_from_filename(stem: str, data: dict[str, Any]) -> tuple[str, str]:
    stem = strip_timestamp_suffix(stem)
    candidate_tasks = sorted(
        set(data.get("results", {}).keys()) | set(data.get("groups", {}).keys()),
        key=len,
        reverse=True,
    )
    for task_name in candidate_tasks:
        prefix = f"{task_name}-"
        if stem.startswith(prefix):
            return task_name, stem[len(prefix):]
        if stem == task_name:
            return task_name, "default"
    if "-" in stem:
        task_name, setting = stem.split("-", 1)
        return task_name, setting
    return stem, "default"


def iter_length_scores(task_payload: dict[str, Any]) -> dict[int, float]:
    out: dict[int, float] = {}
    for key, value in task_payload.items():
        m = _LEN_METRIC_RE.match(key)
        if not m:
            continue
        seq_len = int(m.group(1))
        if isinstance(value, (int, float)):
            out[seq_len] = float(value)
    return out


def get_model_label(model_dir_name: str, data: dict[str, Any]) -> str:
    model_args = data.get("config", {}).get("model_args", {})
    pretrained = model_args.get("pretrained") or data.get("model_name")
    if isinstance(pretrained, str) and pretrained:
        return Path(pretrained).name
    return model_dir_name


def collect_file_rows(model_dir_name: str, jpath: Path) -> tuple[list[dict[str, Any]], dict[str, Any] | None]:
    data = json.loads(jpath.read_text(encoding="utf-8"))
    top_task, setting = parse_task_setting_from_filename(jpath.stem, data)
    model_label = get_model_label(model_dir_name, data)
    model_args = data.get("config", {}).get("model_args", {})

    base_meta = {
        "model": model_dir_name,
        "model_label": model_label,
        "setting": setting,
        "source_file": jpath.name,
        "top_task": top_task,
        "xsa_intervention_site": model_args.get("xsa_intervention_site"),
        "xsa_forward_op": model_args.get("xsa_forward_op"),
        "xsa_forward_alpha": model_args.get("xsa_forward_alpha"),
        "max_length": model_args.get("max_length"),
    }

    results = data.get("results", {})
    subtasks = data.get("group_subtasks", {}).get("ruler")
    if not subtasks:
        subtasks = [k for k in results.keys() if k != "ruler"]

    long_rows: list[dict[str, Any]] = []
    per_len_all: dict[int, list[float]] = {seq: [] for seq in SEQ_LENGTH_ORDER}

    for task_name in subtasks:
        payload = results.get(task_name)
        if not isinstance(payload, dict):
            continue
        scores = iter_length_scores(payload)
        category = categorize_task(task_name)
        for seq_len, score in scores.items():
            long_rows.append(
                {
                    **base_meta,
                    "category": category,
                    "task": task_name,
                    "seq_len": seq_len,
                    "score": round(score * 100, 4),
                }
            )
            per_len_all.setdefault(seq_len, []).append(score)

    overall_payload = results.get("ruler") if isinstance(results.get("ruler"), dict) else None
    overall_scores = iter_length_scores(overall_payload) if overall_payload else {}
    if not overall_scores:
        overall_scores = {
            seq_len: sum(vals) / len(vals)
            for seq_len, vals in per_len_all.items()
            if vals
        }

    overall_row = None
    if overall_scores:
        overall_row = {
            **base_meta,
            "category": "overall",
            "task": "ruler",
        }
        for seq_len, score in overall_scores.items():
            overall_row[f"len_{seq_len}"] = round(score * 100, 4)

    return long_rows, overall_row


def make_wide_table(
    long_df: pd.DataFrame,
    group_cols: list[str],
) -> pd.DataFrame:
    if long_df.empty:
        return pd.DataFrame()
    wide = (
        long_df.pivot_table(
            index=group_cols,
            columns="seq_len",
            values="score",
            aggfunc="mean",
        )
        .reset_index()
    )
    wide.columns = [
        f"len_{int(col)}" if isinstance(col, (int, float)) else str(col)
        for col in wide.columns
    ]
    len_cols = [f"len_{seq}" for seq in SEQ_LENGTH_ORDER if f"len_{seq}" in wide.columns]
    if len_cols:
        wide["avg"] = wide[len_cols].mean(axis=1).round(4)
    return wide


def make_one_table(category_wide: pd.DataFrame, overall_df: pd.DataFrame) -> pd.DataFrame:
    key_cols = ["model_label", "setting"]
    row_map: dict[tuple[str, str], dict[tuple[str, str], Any]] = {}

    for _, row in category_wide.iterrows():
        key = (str(row["model_label"]), str(row["setting"]))
        bucket = row_map.setdefault(key, {})
        category = str(row["category"])
        for seq_len in SEQ_LENGTH_ORDER:
            col = f"len_{seq_len}"
            if col in row and pd.notna(row[col]):
                bucket[(category, f"{seq_len // 1024}k")] = row[col]
        if "avg" in row and pd.notna(row["avg"]):
            bucket[(category, "avg")] = row["avg"]

    for _, row in overall_df.iterrows():
        key = (str(row["model_label"]), str(row["setting"]))
        bucket = row_map.setdefault(key, {})
        for seq_len in SEQ_LENGTH_ORDER:
            col = f"len_{seq_len}"
            if col in row and pd.notna(row[col]):
                bucket[("overall", f"{seq_len // 1024}k")] = row[col]
        if "avg" in row and pd.notna(row["avg"]):
            bucket[("overall", "avg")] = row["avg"]

    if not row_map:
        return pd.DataFrame()

    ordered_cols: list[tuple[str, str]] = [
        ("meta", "model"),
        ("meta", "setting"),
    ]
    for category in ["niah", "qa", "vt", "cwe", "fwe", "overall"]:
        for seq_len in SEQ_LENGTH_ORDER:
            ordered_cols.append((category, f"{seq_len // 1024}k"))
        ordered_cols.append((category, "avg"))

    records: list[list[Any]] = []
    for (model_label, setting), value_map in sorted(row_map.items(), key=lambda x: x[0]):
        row_values: list[Any] = [model_label, setting]
        for col in ordered_cols[2:]:
            row_values.append(value_map.get(col))
        records.append(row_values)

    one_table = pd.DataFrame(records, columns=pd.MultiIndex.from_tuples(ordered_cols))
    return one_table


def write_grouped_excel(df: pd.DataFrame, path: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except Exception:
        print("[WARN] openpyxl is not installed; skip grouped xlsx export.")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "ruler_summary"

    if df.empty:
        wb.save(path)
        return

    cols = list(df.columns)
    top_headers = [col[0] if isinstance(col, tuple) else str(col) for col in cols]
    sub_headers = [col[1] if isinstance(col, tuple) else "" for col in cols]

    center = Alignment(horizontal="center", vertical="center")
    bold = Font(bold=True)

    for idx, (top, sub) in enumerate(zip(top_headers, sub_headers), start=1):
        if top == "meta":
            ws.cell(row=1, column=idx, value=sub)
            ws.merge_cells(start_row=1, start_column=idx, end_row=2, end_column=idx)
        else:
            ws.cell(row=1, column=idx, value=top)
            ws.cell(row=2, column=idx, value=sub)

    start = 3
    for row_values in df.itertuples(index=False, name=None):
        for col_idx, value in enumerate(row_values, start=1):
            ws.cell(row=start, column=col_idx, value=value)
        start += 1

    col_idx = 3
    while col_idx <= len(cols):
        top = top_headers[col_idx - 1]
        start_col = col_idx
        while col_idx <= len(cols) and top_headers[col_idx - 1] == top:
            col_idx += 1
        end_col = col_idx - 1
        if end_col > start_col:
            ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)

    for row in ws.iter_rows(min_row=1, max_row=2):
        for cell in row:
            cell.alignment = center
            cell.font = bold

    ws.freeze_panes = "C3"

    for col_idx in range(1, len(cols) + 1):
        if col_idx == 1:
            ws.column_dimensions["A"].width = 22
        elif col_idx == 2:
            ws.column_dimensions["B"].width = 28
        else:
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = 10

    wb.save(path)


def main() -> None:
    harness_dir = Path(__file__).resolve().parents[1]
    input_root = Path(
        os.environ.get("INPUT_ROOT", harness_dir / "outputs" / "xsa_lm_eval_ruler")
    )
    output_dir = Path(
        os.environ.get("OUTPUT_DIR", harness_dir / "outputs" / "xsa_lm_eval_ruler_summary")
    )

    if not input_root.exists():
        raise SystemExit(f"Input root missing: {input_root}")

    long_rows: list[dict[str, Any]] = []
    overall_rows: list[dict[str, Any]] = []

    for model_dir in sorted(input_root.iterdir()):
        if not model_dir.is_dir():
            continue
        for jpath in sorted(model_dir.glob("*.json")):
            try:
                file_rows, overall_row = collect_file_rows(model_dir.name, jpath)
            except Exception as exc:
                print(f"[ERROR] {jpath}: {exc}")
                continue
            long_rows.extend(file_rows)
            if overall_row is not None:
                overall_rows.append(overall_row)

    if not long_rows:
        raise SystemExit("No RULER rows collected.")

    long_df = pd.DataFrame(long_rows)
    long_df["category_order"] = long_df["category"].map(lambda x: CATEGORY_ORDER.get(str(x), 999))
    long_df = long_df.sort_values(
        ["model_label", "model", "setting", "category_order", "task", "seq_len"]
    ).drop(columns=["category_order"]).reset_index(drop=True)

    task_wide = make_wide_table(
        long_df,
        [
            "model_label",
            "model",
            "setting",
            "xsa_intervention_site",
            "xsa_forward_op",
            "xsa_forward_alpha",
            "category",
            "task",
        ],
    )
    if not task_wide.empty:
        task_wide["category_order"] = task_wide["category"].map(lambda x: CATEGORY_ORDER.get(str(x), 999))
        task_wide = task_wide.sort_values(
            ["model_label", "model", "setting", "category_order", "task"]
        ).drop(columns=["category_order"]).reset_index(drop=True)

    category_wide = make_wide_table(
        long_df,
        [
            "model_label",
            "model",
            "setting",
            "xsa_intervention_site",
            "xsa_forward_op",
            "xsa_forward_alpha",
            "category",
        ],
    )
    if not category_wide.empty:
        category_wide["category_order"] = category_wide["category"].map(lambda x: CATEGORY_ORDER.get(str(x), 999))
        category_wide = category_wide.sort_values(
            ["model_label", "model", "setting", "category_order"]
        ).drop(columns=["category_order"]).reset_index(drop=True)

    overall_df = pd.DataFrame(overall_rows)
    if not overall_df.empty:
        len_cols = [f"len_{seq}" for seq in SEQ_LENGTH_ORDER if f"len_{seq}" in overall_df.columns]
        if len_cols:
            overall_df["avg"] = overall_df[len_cols].mean(axis=1).round(4)
        overall_df = overall_df.sort_values(
            ["model_label", "model", "setting"]
        ).reset_index(drop=True)

    one_table = make_one_table(category_wide, overall_df)

    output_dir.mkdir(parents=True, exist_ok=True)
    one_table_path = output_dir / "ruler_summary_one_table.csv"
    one_table_xlsx_path = output_dir / "ruler_summary_one_table.xlsx"
    long_path = output_dir / "ruler_results_long.csv"
    task_path = output_dir / "ruler_task_by_length.csv"
    category_path = output_dir / "ruler_category_by_length.csv"
    overall_path = output_dir / "ruler_overall_by_length.csv"

    one_table.to_csv(one_table_path, index=False)
    write_grouped_excel(one_table, one_table_xlsx_path)
    long_df.to_csv(long_path, index=False)
    task_wide.to_csv(task_path, index=False)
    category_wide.to_csv(category_path, index=False)
    overall_df.to_csv(overall_path, index=False)

    print(f"[OK] one-table     -> {one_table_path}")
    print(f"[OK] one-table xlsx-> {one_table_xlsx_path}")
    print(f"[OK] long rows     -> {long_path}")
    print(f"[OK] task table    -> {task_path}")
    print(f"[OK] category table-> {category_path}")
    print(f"[OK] overall table -> {overall_path}")
    print()
    if not one_table.empty:
        print(one_table.to_string(index=False))


if __name__ == "__main__":
    main()
